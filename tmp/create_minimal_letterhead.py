from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/Minimal_Letterhead.docx")
NAVY = RGBColor(30, 52, 75)
SLATE = RGBColor(91, 101, 112)
LIGHT = RGBColor(220, 225, 230)


def set_font(run, name="Aptos", size=11, color=None, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def set_paragraph_border(paragraph, color="DCE1E6", size="6", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_body_paragraph(doc, text=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    if text:
        set_font(p.add_run(text), size=11)
    return p


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    # Minimal letterhead header: company identity at left, contact line at right.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), WD_TAB_ALIGNMENT.RIGHT)
    set_font(hp.add_run("YOUR COMPANY"), size=14, color=NAVY, bold=True)
    hp.add_run("\t")
    set_font(hp.add_run("www.yourcompany.com  |  hello@yourcompany.com  |  (000) 000-0000"), size=8.5, color=SLATE)
    rule = header.add_paragraph()
    rule.paragraph_format.space_before = Pt(1)
    rule.paragraph_format.space_after = Pt(0)
    set_paragraph_border(rule)

    # Letter content, intentionally editable with understated placeholders.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run("[MONTH DAY, YEAR]"), size=10.5, color=SLATE)

    for line in ("[RECIPIENT NAME]", "[TITLE / ORGANIZATION]", "[ADDRESS LINE 1]", "[CITY, STATE ZIP]"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(line), size=10.5, color=SLATE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run("Dear [RECIPIENT NAME],"), size=11)

    add_body_paragraph(doc, "Use this space for your letter. The clean layout keeps attention on the message while the header carries your organization details.")
    add_body_paragraph(doc, "Keep paragraphs concise and replace every bracketed placeholder before sending. Add or remove paragraphs as needed.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(24)
    set_font(p.add_run("Sincerely,"), size=11)

    for line, size, color, bold in (("[YOUR NAME]", 11, NAVY, True), ("[TITLE]", 10.5, SLATE, False)):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(line), size=size, color=color, bold=bold)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(4)
    set_font(fp.add_run("YOUR COMPANY  |  123 BUSINESS STREET  |  CITY, STATE 00000"), size=8.5, color=SLATE)

    doc.core_properties.title = "Minimal Letterhead"
    doc.core_properties.subject = "Editable minimal business letterhead template"
    doc.core_properties.author = ""
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
