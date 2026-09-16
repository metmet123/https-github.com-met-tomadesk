import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PyQt6.QtCore import QRect, Qt
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QMessageBox, QStyleOptionViewItem

from alert_notes.editor_shortcut_settings import EditorShortcutSettingsDialog
from alert_notes.memo_list import CenteredCheckDelegate
from alert_notes.panel import AlertNotesPanel
from alert_notes.sqlite_store import NoteReminderStore
from qt_test_support import close_alert_panel
from main_window import TrashDialog


def check_rect(panel, item) -> QRect:
    """0번 열(선택 체크) 칸의 자리.

    트리의 `visualItemRect` 는 줄 전체를 주므로, 칸 하나만 따로 잡아야 한다.
    """
    row = panel.table.visualItemRect(item)
    return QRect(
        panel.table.header().sectionViewportPosition(0), row.y(),
        panel.table.columnWidth(0), row.height(),
    )


class MemoManagementQtTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = NoteReminderStore(self.root / "notes.db", default_title="새 메모")
        self.panel = AlertNotesPanel(self.store)

    def tearDown(self):
        close_alert_panel(self.panel, self.app)
        self.store.close()
        self.temp.cleanup()

    def test_blank_and_option_only_drafts_do_not_create_notes(self):
        self.panel.editor.postit_check.setChecked(True)
        QTest.qWait(450)
        self.assertEqual(self.store.notes(), [])

        self.panel.editor.title_edit.setText("잠시 입력")
        self.panel.editor.title_edit.clear()
        QTest.qWait(450)
        self.assertEqual(self.store.notes(), [])

    def test_title_only_auto_save_creates_exactly_one_note(self):
        self.panel.editor.title_edit.setText("제목만 작성")
        QTest.qWait(500)
        rows = self.store.notes()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["title"], "제목만 작성")

        self.panel.editor.title_edit.setText("제목 수정")
        QTest.qWait(500)
        self.assertEqual(len(self.store.notes()), 1)
        self.assertEqual(self.store.notes()[0]["title"], "제목 수정")

    def test_body_only_manual_save_uses_default_title_and_keeps_full_body(self):
        body = "본문만 작성합니다.\n두 번째 줄입니다."
        self.panel.editor.content_edit.setPlainText(body)
        self.panel.editor._manual_save()
        rows = self.store.notes()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["title"], "새 메모")
        self.assertEqual(self.panel.editor.plain_content(), body)
        self.assertEqual(self.panel.editor.note_id, int(rows[0]["id"]))

    def test_list_checkboxes_line_up_with_the_header_checkbox(self):
        self.store.create_note("첫 메모", "본문")
        self.store.create_note("둘째 메모", "본문")
        self.panel.refresh()
        panel = self.panel.list_panel
        panel.resize(700, 420)
        self.app.processEvents()
        panel.table_header._position_checkbox()

        option = QStyleOptionViewItem()
        option.rect = check_rect(panel, panel.table.topLevelItem(0))
        row_centre = CenteredCheckDelegate.indicator_rect(option).center().x()
        column_centre = option.rect.center().x()
        self.assertLessEqual(
            abs(row_centre - column_centre), 1, "목록 체크가 칸 한가운데에 있지 않습니다",
        )

        header = panel.table_header
        header_centre = header._checkbox.x() + header._indicator_centre().x()
        self.assertLessEqual(
            abs(header_centre - column_centre), 2,
            "머리글 체크와 목록 체크가 어긋납니다",
        )

    def test_clicking_the_centred_checkbox_still_checks_the_row(self):
        note_id = self.store.create_note("첫 메모", "본문")
        self.panel.refresh()
        panel = self.panel.list_panel
        panel.resize(700, 420)
        self.app.processEvents()
        option = QStyleOptionViewItem()
        option.rect = check_rect(panel, panel.table.topLevelItem(0))
        QTest.mouseClick(
            panel.table.viewport(), Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
            CenteredCheckDelegate.indicator_rect(option).center(),
        )
        self.assertEqual(panel.checked_ids(), [int(note_id)])

    def test_search_is_cleared_when_draft_is_created(self):
        self.panel.list_panel.search.setText("검색 결과 없음")
        self.assertEqual(self.panel.list_panel.row_count(), 0)
        self.panel.editor.content_edit.setPlainText("검색 밖에서 바로 작성")
        self.panel.editor._manual_save()
        self.assertEqual(self.panel.list_panel.search.text(), "")
        self.assertEqual(self.panel.list_panel.row_count(), 1)

    def test_save_failure_keeps_one_bound_note_and_manual_retry_succeeds(self):
        real_update = self.store.update_note
        attempts = 0

        def flaky_update(note_id, **values):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise OSError("임시 저장 실패")
            return real_update(note_id, **values)

        self.panel.editor.title_edit.setText("재시도 메모")
        with (
            patch.object(self.store, "update_note", side_effect=flaky_update),
            patch("alert_notes.panel.QMessageBox.warning"),
        ):
            self.panel.editor._manual_save()
            self.assertTrue(self.panel.editor.saved_status.text().endswith("저장 실패"))
            self.panel.editor.content_edit.setPlainText("재시도 본문")
            self.panel.editor._manual_save()
        self.assertEqual(len(self.store.notes()), 1)
        self.assertEqual(self.panel.editor.note_id, int(self.store.notes()[0]["id"]))
        self.assertTrue(self.panel.editor.saved_status.text().endswith("저장됨 · 직접 저장"))
        self.assertIn("재시도 본문", self.panel.editor.plain_content())

    def test_memo_table_sort_check_preview_and_excel(self):
        older = self.store.create_note("이전 메모", "<p>이전 본문</p>")
        newer = self.store.create_note("최근 메모", "<p>첫 줄<br>둘째 줄 전체 본문</p>")
        self.store.update_note(newer, postit=True)
        self.store.conn.execute("UPDATE notes SET updated_at='202608091000' WHERE id=?", (older,))
        self.store.conn.execute("UPDATE notes SET updated_at='202608101230' WHERE id=?", (newer,))
        self.store.conn.commit()
        self.panel.refresh()

        table = self.panel.list_panel.table
        self.assertEqual(
            [table.headerItem().text(i) for i in range(table.columnCount())],
            ["", "제목", "카테고리", "수정일"],
        )
        self.assertFalse(hasattr(self.panel.list_panel, "select_all_button"))
        self.panel.list_panel.table_header.toggle_check_state()
        self.assertEqual(len(self.panel.list_panel.checked_ids()), self.panel.list_panel.row_count())
        table.topLevelItem(0).setCheckState(0, Qt.CheckState.Unchecked)
        self.assertEqual(
            self.panel.list_panel.table_header.check_state(),
            Qt.CheckState.PartiallyChecked,
        )
        # 표시 열이 사라진 자리를 제목 앞 표식이 대신한다.
        self.assertTrue(table.topLevelItem(0).text(1).endswith("최근 메모"))
        self.assertTrue(table.topLevelItem(0).text(1).startswith("📌"))
        self.assertEqual(table.topLevelItem(0).text(2), "—")
        self.assertEqual(table.topLevelItem(0).toolTip(2), "미지정")
        self.assertIn("둘째 줄 전체 본문", table.topLevelItem(0).toolTip(1))

        self.panel.list_panel._set_all_checked(False)
        table.topLevelItem(0).setCheckState(0, Qt.CheckState.Checked)
        output = self.root / "memos.xlsx"
        with (
            patch("alert_notes.panel.QFileDialog.getSaveFileName", return_value=(str(output), "")),
            patch("alert_notes.panel.QMessageBox.information"),
        ):
            self.panel.export_memos()
        workbook = load_workbook(output)
        sheet = workbook["메모목록"]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 3).value, "최근 메모")
        self.assertEqual(sheet.cell(2, 4).value, "미지정")
        self.assertTrue(sheet.cell(2, 5).value)
        self.assertEqual(sheet.cell(2, 6).value, "2026-08-10 12:30")
        self.assertIn("둘째 줄 전체 본문", sheet.cell(2, 7).value)

        with (
            patch("alert_notes.panel.QFileDialog.getSaveFileName", return_value=(str(self.root / "error.xlsx"), "")),
            patch("alert_notes.panel.export_table_xlsx", side_effect=RuntimeError("openpyxl 오류")),
            patch("alert_notes.panel.QMessageBox.warning") as warning,
        ):
            self.panel.export_memos()
        warning.assert_called_once()
        self.assertEqual(len(self.store.notes()), 2)

    def test_selected_note_delete_cascades_reminder(self):
        note_id = self.store.create_note("삭제 대상", "본문")
        self.store.add_reminder(note_id, "209901011200", "알림")
        self.panel.refresh()
        self.panel.list_panel.table.topLevelItem(0).setCheckState(0, Qt.CheckState.Checked)
        with patch(
            "alert_notes.panel.QMessageBox.question",
            return_value=QMessageBox.StandardButton.Yes,
        ):
            self.panel.delete_selected_notes()
        self.assertIsNone(self.store.note(note_id))
        self.assertEqual(self.store.pending_reminders(), [])

    def test_pending_and_history_excel_checked_first(self):
        first = self.store.create_note("첫 알림", "내용")
        second = self.store.create_note("둘째 알림", "내용")
        first_reminder = self.store.add_reminder(first, "209901011200", "첫 메모")
        self.store.add_reminder(second, "209901021200", "둘째 메모")
        history_source = self.store.add_reminder(first, "209901031200", "처리 메모")
        self.store.complete_reminder(history_source)
        self.panel.reminder_history.refresh()

        reminder_panel = self.panel.reminder_history
        for row in range(reminder_panel.pending_table.rowCount()):
            if int(reminder_panel.pending_table.item(row, 1).text()) == first_reminder:
                reminder_panel.pending_table.item(row, 0).setCheckState(Qt.CheckState.Checked)
        pending_path = self.root / "pending.xlsx"
        with (
            patch("alert_notes.reminder_history.QFileDialog.getSaveFileName", return_value=(str(pending_path), "")),
            patch("alert_notes.reminder_history.QMessageBox.information"),
        ):
            reminder_panel._export_pending()
        pending_sheet = load_workbook(pending_path)["예정알림"]
        self.assertEqual(pending_sheet.max_row, 2)
        self.assertEqual(pending_sheet.cell(2, 1).value, first_reminder)

        history_path = self.root / "history.xlsx"
        with (
            patch("alert_notes.reminder_history.QFileDialog.getSaveFileName", return_value=(str(history_path), "")),
            patch("alert_notes.reminder_history.QMessageBox.information"),
        ):
            reminder_panel._export_history()
        history_sheet = load_workbook(history_path)["처리내역"]
        self.assertEqual(history_sheet.max_row, 2)
        self.assertEqual(history_sheet.cell(2, 4).value, "처리 메모")

    def test_standalone_window_is_native_resizable_reusable_and_flushes(self):
        note_id = self.store.create_note("독립창", "원본")
        self.panel.open_standalone_note(note_id)
        window = self.panel.standalone_window
        self.assertIsNotNone(window)
        self.assertTrue(window.windowFlags() & Qt.WindowType.Window)
        self.assertEqual((window.minimumWidth(), window.minimumHeight()), (480, 640))
        self.assertEqual(window.DEFAULT_WIDTH, 600)
        self.assertTrue(window.isVisible())

        window.editor.title_edit.setText("독립창 수정")
        window.close()
        self.app.processEvents()
        self.assertFalse(window.isVisible())
        self.assertEqual(self.store.note(note_id)["title"], "독립창 수정")
        self.assertTrue(self.store.setting(window.GEOMETRY_KEY, ""))

        self.panel.open_standalone_note(note_id)
        self.assertIs(self.panel.standalone_window, window)

    def test_trash_header_selects_and_restores_checked_items(self):
        restored = []
        items = [
            {
                "kind": kind, "id": index, "title": f"항목 {index}",
                "deleted_at": f"20260828090{index}",
                "restore": lambda value=index: restored.append(value),
            }
            for index, kind in enumerate(("단축키 작업", "메모", "일정"), start=1)
        ]
        dialog = TrashDialog(items)
        dialog.resize(680, 420)
        dialog.show()
        self.app.processEvents()
        self.assertEqual(
            [dialog.table.horizontalHeaderItem(i).text() for i in range(4)],
            ["", "종류", "이름", "삭제 시각"],
        )
        self.assertLessEqual(abs(sum(dialog.table.columnWidth(i) for i in range(4)) - dialog.table.viewport().width()), 1)
        dialog.table_header.toggle_check_state()
        self.assertTrue(all(
            dialog.table.item(row, 0).checkState() == Qt.CheckState.Checked
            for row in range(dialog.table.rowCount())
        ))
        dialog.table.item(0, 0).setCheckState(Qt.CheckState.Unchecked)
        self.assertEqual(dialog.table_header.check_state(), Qt.CheckState.PartiallyChecked)
        dialog._restore_selected()
        self.assertEqual(sorted(restored), [2, 3])
        self.assertEqual(dialog.table.rowCount(), 1)
        dialog.table.selectRow(0)
        dialog._restore_selected()
        self.assertEqual(sorted(restored), [1, 2, 3])
        dialog.close()

    def test_note_shortcut_settings_shows_one_success_message(self):
        dialog = EditorShortcutSettingsDialog(self.store)
        with (
            patch("alert_notes.editor_shortcut_settings.QMessageBox.information") as information,
            patch("alert_notes.editor_shortcut_settings.QMessageBox.warning") as warning,
        ):
            dialog._save()
        warning.assert_not_called()
        information.assert_called_once()
        self.assertEqual(information.call_args.args[2], "설정이 저장되었습니다.")

    def test_invalid_note_shortcut_settings_do_not_show_success(self):
        dialog = EditorShortcutSettingsDialog(self.store)
        dialog.always_top.setText("지원안함")
        with (
            patch("alert_notes.editor_shortcut_settings.QMessageBox.information") as information,
            patch("alert_notes.editor_shortcut_settings.QMessageBox.warning") as warning,
        ):
            dialog._save()
        warning.assert_called_once()
        information.assert_not_called()


if __name__ == "__main__":
    unittest.main()
