import json
from pathlib import Path

from foreground_app import normalize_app_list, persisted_app_list
from hotkey_parser import parse_hotkey


SHEET_NAME = "단축키목록"
BASE_HEADERS = ["ID", "활성", "이름", "단축키", "작업유형", "문구", "입력후Enter", "URL", "경로", "매크로JSON"]
LAYOUT_HEADER = "창배치JSON"
EXCLUDED_APPS_HEADER = "제외프로그램"
RESTORE_MINIMIZED_HEADER = "최소화창복원"
HEADERS = [*BASE_HEADERS, LAYOUT_HEADER, RESTORE_MINIMIZED_HEADER, EXCLUDED_APPS_HEADER]
ACTION_TYPES = {"text", "url", "path", "macro", "layout"}


class ExcelImportError(ValueError):
    pass


def export_actions_xlsx(actions, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for action in actions:
        ws.append(action_to_excel_row(action))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _fit_columns(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_actions_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    missing = [header for header in BASE_HEADERS if header not in headers]
    if missing:
        raise ExcelImportError(f"필수 컬럼이 없습니다: {', '.join(missing)}")
    index = {header: headers.index(header) for header in HEADERS if header in headers}
    actions: list[dict] = []
    used_ids: set[int] = set()
    used_hotkeys: set[str] = set()
    errors: list[str] = []
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if _blank_row(values):
            continue
        try:
            action = excel_row_to_action(values, index)
            _validate_unique(action, used_ids, used_hotkeys)
            actions.append(action)
        except Exception as exc:
            errors.append(f"{row_no}행: {exc}")
    if errors:
        raise ExcelImportError("\n".join(errors[:20]))
    return actions


def action_to_excel_row(action) -> list:
    payload = _payload(action)
    action_type = _value(action, "action_type")
    return [
        _value(action, "id"),
        "예" if int(_value(action, "active") or 0) else "아니오",
        _value(action, "name"),
        _value(action, "hotkey"),
        action_type,
        payload.get("text", "") if action_type == "text" else "",
        _bool_label(payload.get("press_enter", True)) if action_type == "text" else "",
        payload.get("url", "") if action_type == "url" else "",
        payload.get("path", "") if action_type == "path" else "",
        json.dumps(payload, ensure_ascii=False, indent=2) if action_type == "macro" else "",
        json.dumps(payload, ensure_ascii=False, indent=2) if action_type == "layout" else "",
        _bool_label(payload.get("restore_if_minimized", False)) if action_type == "path" else "",
        json.dumps(persisted_app_list(payload.get("excluded_apps", [])), ensure_ascii=False)
        if action_type != "layout" else "",
    ]


def excel_row_to_action(values, index: dict[str, int]) -> dict:
    action_id = _parse_id(_cell(values, index, "ID"))
    active = parse_bool(_cell(values, index, "활성"), default=True)
    name = str(_cell(values, index, "이름") or "").strip() or "작업"
    hotkey = parse_hotkey(str(_cell(values, index, "단축키") or "").strip()).text
    action_type = str(_cell(values, index, "작업유형") or "").strip().lower()
    if action_type not in ACTION_TYPES:
        raise ExcelImportError("작업유형은 text/url/path/macro/layout 중 하나여야 합니다.")
    payload = _payload_from_excel(action_type, values, index)
    if action_type == "path":
        payload["restore_if_minimized"] = parse_bool(
            _cell_optional(values, index, RESTORE_MINIMIZED_HEADER),
            default=False,
        )
    if action_type == "layout":
        return {"id": action_id, "name": name, "hotkey": hotkey, "action_type": action_type,
                "payload": payload, "active": active}
    excluded_text = str(_cell_optional(values, index, EXCLUDED_APPS_HEADER) or "").strip()
    if excluded_text:
        try:
            excluded_apps = json.loads(excluded_text)
        except json.JSONDecodeError as exc:
            raise ExcelImportError(f"제외프로그램 형식이 올바르지 않습니다: {exc}") from exc
        if not isinstance(excluded_apps, list):
            raise ExcelImportError("제외프로그램은 JSON 배열이어야 합니다.")
        payload["excluded_apps"] = normalize_app_list(excluded_apps)
    else:
        payload["excluded_apps"] = normalize_app_list(payload.get("excluded_apps", []))
    return {"id": action_id, "name": name, "hotkey": hotkey, "action_type": action_type,
            "payload": payload, "active": active}


def parse_bool(value, default: bool = False) -> bool:
    if value is None or str(value).strip() == "":
        return default
    text = str(value).strip().upper()
    if text in {"1", "Y", "YES", "TRUE", "T", "예", "네"}:
        return True
    if text in {"0", "N", "NO", "FALSE", "F", "아니오", "아니요"}:
        return False
    raise ExcelImportError(f"boolean 값이 아닙니다: {value}")


def _payload_from_excel(action_type: str, values, index: dict[str, int]) -> dict:
    if action_type == "text":
        text = str(_cell(values, index, "문구") or "")
        if not text:
            raise ExcelImportError("text 작업은 문구가 필요합니다.")
        return {"text": text, "press_enter": parse_bool(_cell(values, index, "입력후Enter"), default=True)}
    if action_type == "url":
        url = str(_cell(values, index, "URL") or "").strip()
        if not url:
            raise ExcelImportError("url 작업은 URL이 필요합니다.")
        return {"url": url}
    if action_type == "path":
        target = str(_cell(values, index, "경로") or "").strip()
        if not target:
            raise ExcelImportError("path 작업은 경로가 필요합니다.")
        return {"path": target}
    if action_type == "layout":
        layout_text = str(_cell_optional(values, index, LAYOUT_HEADER) or "").strip()
        if not layout_text:
            raise ExcelImportError("layout 작업은 창배치JSON이 필요합니다.")
        try:
            payload = json.loads(layout_text)
        except json.JSONDecodeError as exc:
            raise ExcelImportError(f"창배치JSON 형식이 올바르지 않습니다: {exc}") from exc
        if not isinstance(payload, dict) or not isinstance(payload.get("windows"), list):
            raise ExcelImportError("layout 작업은 windows 배열이 있는 JSON 객체가 필요합니다.")
        if not payload["windows"]:
            raise ExcelImportError("layout 작업은 하나 이상의 창이 필요합니다.")
        return payload
    macro_text = str(_cell(values, index, "매크로JSON") or "").strip()
    try:
        payload = json.loads(macro_text)
    except json.JSONDecodeError as exc:
        raise ExcelImportError(f"매크로JSON 형식이 올바르지 않습니다: {exc}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("steps"), list):
        raise ExcelImportError("macro 작업은 steps 배열이 있는 JSON 객체가 필요합니다.")
    return payload


def _validate_unique(action: dict, used_ids: set[int], used_hotkeys: set[str]) -> None:
    if action["id"] is not None:
        if action["id"] in used_ids:
            raise ExcelImportError(f"ID가 중복되었습니다: {action['id']}")
        used_ids.add(action["id"])
    if action["hotkey"] in used_hotkeys:
        raise ExcelImportError(f"단축키가 중복되었습니다: {action['hotkey']}")
    used_hotkeys.add(action["hotkey"])


def _parse_id(value) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ExcelImportError(f"ID는 숫자여야 합니다: {value}") from exc
    if number <= 0:
        raise ExcelImportError("ID는 1 이상이어야 합니다.")
    return number


def _payload(action) -> dict:
    payload = _value(action, "payload") or "{}"
    if isinstance(payload, dict):
        return payload
    return json.loads(payload)


def _value(action, key: str):
    if isinstance(action, dict):
        return action.get(key)
    return action[key]


def _cell(values, index: dict[str, int], header: str):
    pos = index[header]
    return values[pos] if pos < len(values) else None


def _cell_optional(values, index: dict[str, int], header: str):
    return _cell(values, index, header) if header in index else None


def _blank_row(values) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def _bool_label(value) -> str:
    return "예" if bool(value) else "아니오"


def _fit_columns(ws) -> None:
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 10), 80)
