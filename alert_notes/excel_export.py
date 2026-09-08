from pathlib import Path


def export_table_xlsx(sheet_name: str, headers: list[str], rows: list[list[object]], path: Path) -> Path:
    """Write one filtered/selected alert-note table to a standalone workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    output = Path(path)
    if output.suffix.casefold() != ".xlsx":
        output = output.with_suffix(".xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    _fit_columns(sheet)
    workbook.save(output)
    return output


def display_datetime(value) -> str:
    text = str(value or "")[:12]
    if len(text) != 12 or not text.isdigit():
        return ""
    return f"{text[:4]}-{text[4:6]}-{text[6:8]} {text[8:10]}:{text[10:12]}"


def _fit_columns(sheet) -> None:
    for cells in sheet.columns:
        maximum = max(len(str(cell.value or "")) for cell in cells)
        sheet.column_dimensions[cells[0].column_letter].width = min(max(maximum + 2, 10), 60)
